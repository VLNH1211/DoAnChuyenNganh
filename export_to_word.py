import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def style_table(table, col_widths, headers, data, header_bg="1B365D", alt_bg="F8FAFC"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        cell = hdr_cells[i]
        cell.width = Inches(col_widths[i])
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=140, bottom=140, left=150, right=150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        is_alt = (r_idx % 2 == 1)
        bg = alt_bg if is_alt else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row_cells[c_idx]
            cell.width = Inches(col_widths[c_idx])
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(val)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10.5)
            if c_idx == 0 or (len(row_data) > 3 and c_idx in [2, 3]):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

def create_word_document():
    doc = docx.Document()

    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.8)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(27, 54, 93)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(43, 87, 154)
        return p

    def add_p(text, bold_prefix="", italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(12)
        r = p.add_run(text)
        r.italic = italic
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        return p

    def add_image_with_caption(img_path, caption_text):
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.paragraph_format.keep_with_next = True
            run = p_img.add_run()
            run.add_picture(img_path, width=Inches(6.2))

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(12)
            r_cap = p_cap.add_run(caption_text)
            r_cap.italic = True
            r_cap.font.name = 'Times New Roman'
            r_cap.font.size = Pt(10.5)
            r_cap.font.color.rgb = RGBColor(71, 85, 105)
        else:
            add_p(f"[LỖI: Không tìm thấy ảnh tại {img_path}]", bold_prefix="CẢNH BÁO: ")

    def add_callout(text, bg="F0FDF4", border_color="10B981"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.3)
        set_cell_background(cell, bg)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    base_dir = os.path.dirname(os.path.abspath(__file__))
    screenshots_dir = os.path.join(base_dir, "screenshots")

    # ---------------------------------------------------------
    # TRANG TIÊU ĐỀ & THÔNG TIN HỌC THUẬT
    # ---------------------------------------------------------
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(10)
    p_inst.paragraph_format.space_after = Pt(2)
    r = p_inst.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ TP. HỒ CHÍ MINH (HUTECH)\nVIỆN CÔNG NGHỆ THÔNG TIN - NGÀNH AN TOÀN THÔNG TIN")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(27, 54, 93)

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(16)
    r_div = p_div.add_run("------------------------------------------------------------------------")
    r_div.font.color.rgb = RGBColor(203, 213, 225)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(8)
    r = p_title.add_run("BỘ MINH CHỨNG HÌNH ẢNH THỰC NGHIỆM ĐỐI CHIẾU\nBEFORE / AFTER (BASELINE V1 VS HARDENED V2)")
    r.bold = True
    r.font.size = Pt(17)
    r.font.color.rgb = RGBColor(185, 28, 28)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    r = p_sub.add_run("ĐỀ TÀI: XÂY DỰNG VÀ KIỂM THỬ CÁC CƠ CHẾ BẢO MẬT CHO ỨNG DỤNG WEB/API DỰA TRÊN XÁC THỰC, PHÂN QUYỀN VÀ KIỂM SOÁT TRUY CẬP")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(27, 54, 93)

    tbl_sv = doc.add_table(rows=5, cols=2)
    tbl_sv.alignment = WD_TABLE_ALIGNMENT.CENTER
    sv_data = [
        ("Sinh viên thực hiện:", "Võ Lê Nhật Hoàng"),
        ("Mã số sinh viên (MSSV):", "2387700017"),
        ("Lớp học phần:", "23DATA1"),
        ("Chuyên ngành:", "An toàn thông tin"),
        ("Phạm vi thực nghiệm:", "4 Kịch bản cốt lõi (KB1–KB4) + 4 Kịch bản mở rộng (KB5–KB8) + OWASP Headers & Audit Log")
    ]
    for idx, (k, v) in enumerate(sv_data):
        c0, c1 = tbl_sv.rows[idx].cells
        c0.width, c1.width = Inches(2.3), Inches(4.0)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, 80, 80, 100, 100)
        set_cell_margins(c1, 80, 80, 100, 100)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(k)
        r0.bold = True
        r0.font.size = Pt(11)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(v)
        if idx == 0:
            r1.bold = True
        r1.font.size = Pt(11)

    doc.add_page_break()

    # ---------------------------------------------------------
    # PHẦN 1: TỔNG QUAN GIAO DIỆN BẢNG ĐIỀU KHIỂN
    # ---------------------------------------------------------
    add_h1("PHẦN 1: TỔNG QUAN GIAO DIỆN BẢNG ĐIỀU KHIỂN THỰC NGHIỆM")
    add_p("Hệ thống thực nghiệm được thiết kế giao diện Web Cyber Security Dashboard với hai phiên bản chạy song song để minh họa trực quan sự khác biệt về mặt an toàn thông tin:")
    add_p(" Phiên bản V1 (http://localhost:5001): Tông màu Đỏ Cảnh báo (Red Team Baseline), tích hợp 8 kịch bản khai thác lỗ hổng (KB1–KB4 cốt lõi & KB5–KB8 mở rộng).", bold_prefix="•")
    add_p(" Phiên bản V2 (http://localhost:5002): Tông màu Xanh Lục An Toàn (Blue Team Hardened), tích hợp đầy đủ 8 cơ chế phòng thủ, OWASP HTTP Security Headers và Audit Logs.", bold_prefix="•")

    add_h2("1.1 Bảng điều khiển Phiên bản V1 (Baseline - Chưa bảo mật)")
    add_image_with_caption(
        os.path.join(screenshots_dir, "01_V1_Dashboard_Overview.png"),
        "Hình 1.1: Giao diện Bảng điều khiển Bản V1 tại http://localhost:5001 với tông đỏ cảnh báo và 2 hàng kịch bản KB1–KB4 & KB5–KB8."
    )

    add_h2("1.2 Bảng điều khiển Phiên bản V2 (Hardened - Đã bảo mật đa lớp)")
    add_image_with_caption(
        os.path.join(screenshots_dir, "02_V2_Dashboard_Overview.png"),
        "Hình 1.2: Giao diện Bảng điều khiển Bản V2 tại http://localhost:5002 với tông xanh an toàn, tích hợp giải mã JWT Claims, phòng thủ KB1–KB8 và Audit Logs."
    )

    # ---------------------------------------------------------
    # PHẦN 2: CẶP MINH CHỨNG KB1
    # ---------------------------------------------------------
    add_h1("PHẦN 2: CẶP MINH CHỨNG KB1 - KIỂM SOÁT XÁC THỰC (AUTHENTICATION CONTROL)")
    add_p("Kiểm tra khả năng ngăn chặn truy cập nặc danh (Anonymous) vào tài nguyên đơn hàng nhạy cảm (GET /api/orders).", bold_prefix="• Mục tiêu kiểm thử: ")
    add_p("Bắt buộc JWT Bearer Token hợp lệ theo chuẩn RFC 7519; nếu không có Token, API phải từ chối ngay lập tức với mã HTTP 401 Unauthorized.", bold_prefix="• Tiêu chuẩn bảo mật: ")

    t_kb1 = doc.add_table(rows=5, cols=3)
    style_table(t_kb1, [1.8, 2.2, 2.3],
                ["Tiêu chí đối chiếu", "Bản V1 (Chưa bảo mật)", "Bản V2 (Đã bảo mật)"],
                [
                    ["Endpoint & Method", "GET /api/orders", "GET /api/orders"],
                    ["Header Authorization", "Không gửi (Nặc danh)", "Không gửi (Nặc danh)"],
                    ["Mã trạng thái HTTP", "200 OK (LỖ HỔNG XÁC THỰC)", "401 Unauthorized (ĐÃ BẢO VỆ)"],
                    ["Dữ liệu phản hồi", "Rò rỉ toàn bộ đơn hàng cho người lạ", "Từ chối: Chưa xác thực hoặc Token không hợp lệ"]
                ])

    add_h2("2.1 [BEFORE] Bản V1 cho phép nặc danh truy xuất dữ liệu (200 OK)")
    add_image_with_caption(
        os.path.join(screenshots_dir, "03_KB1_Before_V1_Unauthenticated_200.png"),
        "Hình 2.1: Lỗ hổng KB1 trên bản V1 - Truy cập nặc danh GET /api/orders trả về 200 OK và làm lộ toàn bộ đơn hàng trong CSDL."
    )

    add_h2("2.2 [AFTER] Bản V2 chặn đứng yêu cầu nặc danh (401 Unauthorized)")
    add_image_with_caption(
        os.path.join(screenshots_dir, "04_KB1_After_V2_Unauthorized_401.png"),
        "Hình 2.2: Cơ chế phòng thủ KB1 trên bản V2 - Chặn đứng yêu cầu nặc danh với mã 401 Unauthorized."
    )

    add_h2("2.3 [AFTER SUCCESS] Bản V2 xác thực thành công JWT Bearer Token của userA")
    add_image_with_caption(
        os.path.join(screenshots_dir, "05_KB1_After_V2_ValidToken_UserA_200.png"),
        "Hình 2.3: Đăng nhập hợp lệ userA trên bản V2 - Nhận JWT Token (có JTI), giải mã Claims và phân tách phạm vi dữ liệu đơn hàng của userA."
    )

    # ---------------------------------------------------------
    # PHẦN 3: CẶP MINH CHỨNG KB2
    # ---------------------------------------------------------
    add_h1("PHẦN 3: CẶP MINH CHỨNG KB2 - PHÂN QUYỀN VAI TRÒ QUẢN TRỊ (RBAC)")
    add_p("Ngăn chặn người dùng có vai trò khách hàng thông thường (Role: NguoiDung) truy cập trái phép vào các chức năng quản trị viên (GET /api/admin/users).", bold_prefix="• Mục tiêu kiểm thử: ")
    add_p("Nguyên tắc đặc quyền tối thiểu (Least Privilege) kết hợp ASP.NET Core Authorization Policy 'AdminOnly'.", bold_prefix="• Nguyên tắc áp dụng: ")

    t_kb2 = doc.add_table(rows=5, cols=3)
    style_table(t_kb2, [1.8, 2.2, 2.3],
                ["Tiêu chí đối chiếu", "Bản V1 (Chưa bảo mật)", "Bản V2 (Đã bảo mật)"],
                [
                    ["Endpoint Quản trị", "GET /api/admin/users", "GET /api/admin/users"],
                    ["Tài khoản thử nghiệm", "userA (Vai trò: NguoiDung)", "userA (Vai trò: NguoiDung)"],
                    ["Mã trạng thái HTTP", "200 OK (LỖ HỔNG BROKEN RBAC)", "403 Forbidden (ĐÃ BẢO VỆ)"],
                    ["Mức độ bảo vệ mật khẩu", "Hiển thị mật khẩu Plaintext thô", "Băm an toàn BCrypt / Ẩn đối với người dùng"]
                ])

    add_h2("3.1 [BEFORE] Bản V1 để lộ API Quản trị & mật khẩu thô cho user thường (200 OK)")
    add_image_with_caption(
        os.path.join(screenshots_dir, "06_KB2_Before_V1_BrokenRBAC_200.png"),
        "Hình 3.1: Lỗ hổng Broken RBAC trên bản V1 - Khách hàng thông thường userA vẫn xem được danh sách Admin và mật khẩu Plaintext."
    )

    add_h2("3.2 [AFTER] Bản V2 chặn đứng yêu cầu vượt quyền của userA (403 Forbidden)")
    add_image_with_caption(
        os.path.join(screenshots_dir, "07_KB2_After_V2_RBAC_Blocked_403.png"),
        "Hình 3.2: Cơ chế phòng thủ RBAC trên bản V2 - Chặn đứng nỗ lực truy cập API Quản trị của userA với mã 403 Forbidden."
    )

    add_h2("3.3 [AFTER SUCCESS] Bản V2 chỉ cho phép Quản trị viên (admin) truy cập hợp lệ (200 OK)")
    add_image_with_caption(
        os.path.join(screenshots_dir, "08_KB2_After_V2_Admin_Access_200.png"),
        "Hình 3.3: Quản trị viên (admin) truy cập hợp lệ trên bản V2 - Mật khẩu được băm BCrypt an toàn, loại trừ hoàn toàn nguy cơ lộ mật khẩu thô."
    )

    # ---------------------------------------------------------
    # PHẦN 4: CẶP MINH CHỨNG KB3
    # ---------------------------------------------------------
    add_h1("PHẦN 4: CẶP MINH CHỨNG KB3 - KIỂM SOÁT ĐỐI TƯỢNG BOLA / IDOR")
    add_p("Ngăn chặn hành vi thao túng ID đơn hàng (id=3) trên URL để xem trộm đơn hàng và thông tin tài chính của khách hàng khác.", bold_prefix="• Mục tiêu kiểm thử: ")
    add_p("Kiểm soát truy cập mức đối tượng (Object-Level Access Control): Đối chiếu trường UserId trong đơn hàng với ClaimTypes.NameIdentifier trích xuất từ Token JWT.", bold_prefix="• Cơ chế bảo vệ: ")

    t_kb3 = doc.add_table(rows=5, cols=3)
    style_table(t_kb3, [1.8, 2.2, 2.3],
                ["Tiêu chí đối chiếu", "Bản V1 (Chưa bảo mật)", "Bản V2 (Đã bảo mật)"],
                [
                    ["Endpoint gọi đến", "GET /api/orders/3", "GET /api/orders/3"],
                    ["Tài khoản thực hiện", "userA (Chủ sở hữu đơn #1, #2, #5)", "userA (Chủ sở hữu đơn #1, #2, #5)"],
                    ["Mã trạng thái HTTP", "200 OK (LỖ HỔNG BOLA / IDOR)", "403 Forbidden (ĐÃ BẢO VỆ)"],
                    ["Hậu quả an toàn", "userA xem trộm đơn 12.5 triệu của userB", "Chặn đứng và từ chối cung cấp dữ liệu"]
                ])

    add_h2("4.1 [BEFORE] Bản V1 cho phép userA xem trộm đơn hàng #3 của userB (200 OK)")
    add_image_with_caption(
        os.path.join(screenshots_dir, "09_KB3_Before_V1_IDOR_Leaked_200.png"),
        "Hình 4.1: Lỗ hổng IDOR trên bản V1 - userA xem trộm đơn hàng #3 trị giá 12.500.000 VNĐ của khách hàng Trần Thị Bích (userB)."
    )

    add_h2("4.2 [AFTER] Bản V2 chặn đứng hành vi IDOR của userA (403 Forbidden)")
    add_image_with_caption(
        os.path.join(screenshots_dir, "10_KB3_After_V2_IDOR_Blocked_403.png"),
        "Hình 4.2: Cơ chế chống BOLA/IDOR trên bản V2 - Chặn đứng truy cập trái phép đơn hàng #3 với mã 403 Forbidden."
    )

    # ---------------------------------------------------------
    # PHẦN 5: CẶP MINH CHỨNG KB4
    # ---------------------------------------------------------
    add_h1("PHẦN 5: CẶP MINH CHỨNG KB4 - KIỂM SOÁT TOÀN VẸN & CHỐNG SQL INJECTION")
    add_p("Ngăn chặn tấn công chèn mã SQL độc hại thông qua chuỗi tìm kiếm đơn hàng (GET /api/orders/search?keyword=...).", bold_prefix="• Mục tiêu kiểm thử: ")
    add_p("Lọc kiểm tra định dạng đầu vào bằng biểu thức chính quy (Regex) và chuyển đổi 100% sang truy vấn tham số hóa LINQ Parameterized Query.", bold_prefix="• Giải pháp phòng thủ trên V2: ")

    t_kb4 = doc.add_table(rows=5, cols=3)
    style_table(t_kb4, [1.8, 2.2, 2.3],
                ["Tiêu chí đối chiếu", "Bản V1 (Chưa bảo mật)", "Bản V2 (Đã bảo mật)"],
                [
                    ["Payload đầu vào", "' OR 1=1 --", "' OR 1=1 --"],
                    ["Kỹ thuật truy vấn SQL", "Ghép chuỗi SQL thô (ExecuteReader)", "Lọc Regex + Parameterized Query EF Core"],
                    ["Mã trạng thái HTTP", "200 OK (LỖ HỔNG SQL INJECTION)", "400 Bad Request (ĐÃ BẢO VỆ)"],
                    ["Tác động hệ thống", "Vô hiệu WHERE, đổ toàn bộ CSDL", "Chặn đứng tiêm mã, CSDL nguyên vẹn an toàn"]
                ])

    add_h2("5.1 [BEFORE] Bản V1 bị khai thác SQL Injection trích xuất toàn bộ dữ liệu (200 OK)")
    add_image_with_caption(
        os.path.join(screenshots_dir, "11_KB4_Before_V1_SQLi_Dumped_200.png"),
        "Hình 5.1: Lỗ hổng SQL Injection trên bản V1 - Chuỗi ' OR 1=1 -- làm thay đổi mệnh đề WHERE, đổ toàn bộ danh sách đơn hàng ra ngoài."
    )

    add_h2("5.2 [AFTER] Bản V2 chặn đứng chuỗi SQL độc hại (400 Bad Request)")
    add_image_with_caption(
        os.path.join(screenshots_dir, "12_KB4_After_V2_SQLi_Blocked_400.png"),
        "Hình 5.2: Cơ chế phòng thủ SQL Injection trên bản V2 - Chặn đứng payload độc hại với mã 400 Bad Request an toàn, không gây crash 500."
    )

    # ---------------------------------------------------------
    # PHẦN 6: BẢNG TỔNG HỢP RUN ALL 4 KỊCH BẢN CỐT LÕI (KB1-KB4)
    # ---------------------------------------------------------
    add_h1("PHẦN 6: MINH CHỨNG BẢNG TỔNG HỢP 4 KỊCH BẢN CỐT LÕI (RUN ALL KB1–KB4)")
    add_p("Chức năng 'Chạy tất cả 4 kịch bản (Run All KB1–KB4)' tự động gửi chuỗi tấn công liên hoàn qua 4 kịch bản cốt lõi theo đề cương nghiên cứu.")

    add_h2("6.1 Kết quả kiểm tra liên hoàn KB1–KB4 trên Bản V1 - Cảnh báo 4/4 Lỗ hổng")
    add_image_with_caption(
        os.path.join(screenshots_dir, "13a_V1_RunAll.png"),
        "Hình 6.1: Kết quả Run All KB1–KB4 trên Bản V1 - Toàn bộ 4 kịch bản chuyển màu ĐỎ (HTTP 200 - LỖ HỔNG)."
    )

    add_h2("6.2 Kết quả kiểm tra liên hoàn KB1–KB4 trên Bản V2 - Xác nhận 4/4 An toàn")
    add_image_with_caption(
        os.path.join(screenshots_dir, "13b_V2_RunAll.png"),
        "Hình 6.2: Kết quả Run All KB1–KB4 trên Bản V2 - Toàn bộ 4 kịch bản chuyển màu XANH LỤC (401, 403, 403, 400 - AN TOÀN)."
    )

    # ---------------------------------------------------------
    # PHẦN 7: MINH CHỨNG 4 KỊCH BẢN PHÒNG THỦ CHIỀU SÂU MỞ RỘNG (KB5–KB8 & OWASP HEADERS)
    # ---------------------------------------------------------
    add_h1("PHẦN 7: MINH CHỨNG CÁC CƠ CHẾ PHÒNG THỦ CHIỀU SÂU MỞ RỘNG (KB5–KB8 & OWASP HEADERS)")
    add_p("Bên cạnh 4 kịch bản cốt lõi trong đề cương, hệ thống mở rộng thêm 4 cơ chế phòng thủ chiều sâu (Defense-in-Depth) nhằm giải quyết các bài toán thực tiễn của hệ thống RESTful API doanh nghiệp:")
    add_p(" Thu hồi JWT Token vào JTI Blacklist ngay khi người dùng đăng xuất (POST /api/auth/logout), ngăn chặn triệt để tấn công phát lại phiên (Token Replay - 401 Unauthorized).", bold_prefix="• KB5 (Token Revocation): ")
    add_p(" Tự động trích xuất UserId từ JWT Claims để chống gán tham số hàng loạt (Mass Assignment), chặn đơn hàng số tiền âm (400 Bad Request) và chặn IDOR khi Sửa/Xóa đơn (403 Forbidden).", bold_prefix="• KB6 (Business Logic & Anti-Mass Assignment): ")
    add_p(" Kiểm tra chữ ký byte đầu tệp thực tế (FF D8 FF cho JPEG, 89 50 4E 47 cho PNG, %PDF cho PDF) và chặn Double Extension để ngăn chặn tải lên Webshell thực thi mã từ xa (400/415).", bold_prefix="• KB7 (Magic Bytes File Upload Validation): ")
    add_p(" Giới hạn tần suất đăng nhập sai và tần suất gọi API bằng thuật toán cửa sổ trượt (Sliding Window), trả về mã chuẩn 429 Too Many Requests kèm Retry-After Header.", bold_prefix="• KB8 (Rate Limiting & Anti-Brute-Force): ")

    t_ext = doc.add_table(rows=5, cols=3)
    style_table(t_ext, [1.8, 2.2, 2.3],
                ["Kịch bản mở rộng", "Bản V1 (Chưa bảo mật)", "Bản V2 (Đã bảo mật)"],
                [
                    ["KB5: Thu hồi Token (Logout)", "200 OK (Token cũ vẫn dùng lại được)", "401 Unauthorized (JTI Blacklist chặn)"],
                    ["KB6: Logic tiền âm & IDOR Xóa", "200 OK (Tạo tiền âm, xóa đơn userB)", "400 Bad Request / 403 Forbidden"],
                    ["KB7: Upload Webshell giả ảnh", "200 OK (Lưu trực tiếp webshell.jpg.aspx)", "400 / 415 (Chặn Magic Bytes giả mạo)"],
                    ["KB8: Spam API liên tục", "200 OK (Không giới hạn tần suất)", "429 Too Many Requests (Retry-After: 5s)"]
                ])

    add_h2("7.1 [BEFORE] Kết quả chạy liên hoàn KB5–KB8 trên Bản V1 - Phát hiện 4/4 Lỗ hổng mở rộng")
    add_image_with_caption(
        os.path.join(screenshots_dir, "15a_V1_RunAll_Extended_KB5_KB8.png"),
        "Hình 7.1: Kết quả Run All KB5–KB8 trên Bản V1 - Bị khai thác Token Replay, chấp nhận đơn hàng tiền âm, cho phép upload file giả mạo và không giới hạn tần suất spam API."
    )

    add_h2("7.2 [AFTER] Kết quả chạy liên hoàn KB5–KB8 trên Bản V2 - Chặn đứng 4/4 cuộc tấn công")
    add_image_with_caption(
        os.path.join(screenshots_dir, "15b_V2_RunAll_Extended_KB5_KB8.png"),
        "Hình 7.2: Kết quả Run All KB5–KB8 trên Bản V2 - Toàn bộ 4 cơ chế phòng thủ chiều sâu (JTI Blacklist 401, Business Logic 400, Magic Bytes 400/415, Rate Limit 429) hoạt động chính xác."
    )

    add_h2("7.3 Kiểm tra trực tiếp bộ 5 OWASP HTTP Security Headers trên Bản V2")
    add_image_with_caption(
        os.path.join(screenshots_dir, "16_V2_OWASP_Security_Headers_LiveCheck.png"),
        "Hình 7.3: Kiểm tra trực tiếp HTTP Response Headers trên Bản V2 xác nhận đầy đủ X-Content-Type-Options: nosniff, X-Frame-Options: DENY, X-XSS-Protection, Referrer-Policy và Content-Security-Policy."
    )

    # ---------------------------------------------------------
    # PHẦN 8: NHẬT KÝ KIỂM TOÁN AN NINH (AUDIT LOGS)
    # ---------------------------------------------------------
    add_h1("PHẦN 8: MINH CHỨNG HỆ THỐNG NHẬT KÝ KIỂM TOÁN AN NINH (SECURITY AUDIT LOGS)")
    add_p("Hệ thống Nhật ký Kiểm toán An ninh trên Bản V2 tự động ghi nhận vết mọi sự kiện an ninh (401, 403, 400, 415, 429, 200, 201) thông qua SecurityAuditMiddleware và lưu trữ trong bảng CSDL AuditLogs phục vụ công tác điều tra số (Digital Forensics) và phản ứng sự cố (Incident Response).")

    add_h2("8.1 Bảng dữ liệu AuditLogs và Thống kê Vi phạm trên Bản V2")
    add_image_with_caption(
        os.path.join(screenshots_dir, "14_V2_Security_Audit_Logs_Table.png"),
        "Hình 8.1: Bảng Nhật ký Kiểm toán An ninh trên Bản V2 hiển thị chi tiết các sự kiện 401, 403, 400, 429 kèm Username, IP, URL và thống kê phát hiện xâm nhập."
    )

    # ---------------------------------------------------------
    # TỔNG KẾT MA TRẬN
    # ---------------------------------------------------------
    add_h1("TỔNG KẾT MA TRẬN KẾT QUẢ THỰC NGHIỆM ĐỒ ÁN (KB1 – KB8)")
    t_sum = doc.add_table(rows=10, cols=5)
    style_table(t_sum, [0.6, 1.8, 1.7, 1.3, 0.9],
                ["STT", "Kịch bản kiểm thử", "Bản V1 (Chưa bảo mật)", "Bản V2 (Đã bảo mật)", "Đánh giá"],
                [
                    ["1", "KB1: Kiểm soát xác thực", "200 OK (Rò rỉ dữ liệu)", "401 Unauthorized (Chặn đứng)", "ĐẠT (TC1)"],
                    ["2", "KB2: Phân quyền vai trò", "200 OK (Lộ Plaintext pass)", "403 Forbidden (Chặn RBAC)", "ĐẠT (TC2)"],
                    ["3", "KB3: Kiểm soát đối tượng", "200 OK (Xem trộm BOLA)", "403 Forbidden (Chặn BOLA)", "ĐẠT (TC3)"],
                    ["4", "KB4: Kiểm soát đầu vào", "200 OK (Dump CSDL qua SQLi)", "400 Bad Request (Chặn Regex)", "ĐẠT (TC4)"],
                    ["5", "KB5: Thu hồi JWT (Logout)", "200 OK (Bị Token Replay)", "401 TokenRevoked (JTI Blacklist)", "ĐẠT (Mở rộng)"],
                    ["6", "KB6: Logic & IDOR Sửa/Xóa", "200 OK (Tiền âm, xóa đơn #3)", "400 / 403 (Bảo vệ nghiệp vụ)", "ĐẠT (Mở rộng)"],
                    ["7", "KB7: Upload Magic Bytes", "200 OK (Chấp nhận Webshell)", "400 / 415 (Chặn chữ ký giả)", "ĐẠT (Mở rộng)"],
                    ["8", "KB8: Rate Limiting & Spam", "200 OK (Không giới hạn)", "429 Too Many Requests", "ĐẠT (Mở rộng)"],
                    ["9", "TC5: OWASP Headers & Audit", "Thiếu Headers / Không ghi vết", "Đủ 5 OWASP Headers & AuditLogs", "ĐẠT (TC5)"]
                ])

    add_callout(
        "KẾT LUẬN THỰC NGHIỆM: Hệ thống WebApi.V2_Secure đã triển khai thành công mô hình phòng thủ đa lớp (4 kịch bản cốt lõi KB1–KB4 theo đề cương kết hợp 4 cơ chế phòng thủ chiều sâu KB5–KB8), vượt qua 100% cả 17 bài kiểm thử tự động, chứng minh tính đúng đắn và giá trị thực tiễn cao của đề tài nghiên cứu.",
        bg="F0FDF4", border_color="10B981"
    )

    out_docs = os.path.join(base_dir, "docs", "BO_MINH_CHUNG_HINH_ANH_BEFORE_AFTER.docx")
    doc.save(out_docs)
    print(f"=== ĐÃ XUẤT FILE WORD THÀNH CÔNG TẠI: {out_docs} ===")

if __name__ == "__main__":
    create_word_document()
